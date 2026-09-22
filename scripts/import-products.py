import json
import re
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image as PILImage

SOURCE = Path("/workspace/scratch/7397f06d92ee/upload")
PUBLIC = Path("/workspace/sites/pedido-central/public/products")
DATA = Path("/workspace/sites/pedido-central/lib/products.json")
PUBLIC.mkdir(parents=True, exist_ok=True)


def number(value):
    if value is None:
        return 0
    match = re.search(r"\d+(?:[.,]\d+)?", str(value))
    return float(match.group(0).replace(",", ".")) if match else 0


def save_images(workbook, sheet, prefix):
    found = []
    for index, image in enumerate(sheet._images, start=1):
        try:
            raw = image._data()
            converted = PILImage.open(BytesIO(raw)).convert("RGBA")
            filename = f"{prefix}-{index}.png"
            converted.save(PUBLIC / filename, "PNG")
            found.append((image.anchor._from.row + 1, f"/products/{filename}"))
        except Exception:
            continue
    return sorted(found)


def nearest_image(images, row):
    valid = [item for item in images if item[0] <= row]
    return valid[-1][1] if valid else (images[0][1] if images else None)


products = []


def add(supplier, description, row, code="", category="Geral", stock=0, order=0, unit="un", image=None):
    products.append({
        "id": f"{supplier.lower().replace(' ', '-')}-{len(products)+1}",
        "supplier": supplier,
        "brand": supplier,
        "code": str(code).strip(),
        "description": str(description).strip(),
        "category": category,
        "unit": unit,
        "stock": stock,
        "suggested": order,
        "image": image,
        "sourceRow": row,
    })


# Force Line
wb = load_workbook(SOURCE / "FORCE LINE (1)(1).xlsx", data_only=False)
ws = wb.worksheets[0]
images = save_images(wb, ws, "force")
for row in range(3, ws.max_row + 1):
    description = ws.cell(row, 2).value
    if not description:
        continue
    text = str(description).strip()
    if "EXTENSÃO" in text or "FILTRO" in text:
        category = "Extensões e filtros"
    elif "MANGUEIRA" in text or "ESGUICHO" in text or "IRRIGADOR" in text:
        category = "Jardim"
    else:
        category = "Acessórios elétricos"
    add("Force Line", text, row, category=category, stock=number(ws.cell(row, 3).value), order=number(ws.cell(row, 4).value), image=nearest_image(images, row))

# Jamaica
wb = load_workbook(SOURCE / "JAMAICA (2)(1).xlsx", data_only=False)
ws = wb.worksheets[0]
category = "Mangueiras Jamaica"
for row in range(3, ws.max_row + 1):
    code = ws.cell(row, 1).value
    measure = ws.cell(row, 2).value
    if code and measure and re.search(r"\d", str(code)):
        add("Jamaica", f"{category} {str(measure).strip()}", row, code=code, category=category, stock=number(ws.cell(row, 3).value), order=number(ws.cell(row, 4).value), unit="un")
    elif code and not measure:
        category = str(code).strip()

# Tubo PU
wb = load_workbook(SOURCE / "TUBO PU 13.01 (2)(1).xlsx", data_only=False)
ws = wb.worksheets[0]
images = save_images(wb, ws, "tubo-pu")
for row in range(2, ws.max_row + 1):
    description = ws.cell(row, 1).value
    if description:
        add("Tubo PU", description, row, category="Tubos pneumáticos", order=number(ws.cell(row, 2).value), unit="m", image=nearest_image(images, row))

# Sucção
wb = load_workbook(SOURCE / "PEDIDO SUCÇÃO 12 (1)(1).xlsx", data_only=False)
ws = wb.worksheets[0]
for row in range(3, ws.max_row + 1):
    for col, color in ((1, "Azul"), (4, "Laranja")):
        measure = ws.cell(row, col).value
        if measure:
            add("Sucção", f"Mangueira de sucção {color} {measure}", row, category="Mangueiras de sucção", order=number(ws.cell(row, col + 1).value), unit="m")

# Bariflex
wb = load_workbook(SOURCE / "BARIFLEX MANG JARDIM BHS(2).xlsx", data_only=False)
ws = wb.worksheets[0]
images = save_images(wb, ws, "bariflex")
for row in range(3, ws.max_row + 1):
    description = ws.cell(row, 2).value
    if description:
        add("Bariflex", description, row, category="Mangueiras de jardim", stock=number(ws.cell(row, 3).value), order=number(ws.cell(row, 4).value), unit="rolo", image=nearest_image(images, row))

DATA.write_text(json.dumps(products, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"{len(products)} produtos, {len(list(PUBLIC.glob('*')))} imagens")
