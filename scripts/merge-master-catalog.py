import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/workspace/scratch/7397f06d92ee/upload/ebd25f18-7cd2-4920-bdee-055d127d7b34.xlsx")
TARGET = ROOT / "lib" / "products.json"


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalized(value):
    text = unicodedata.normalize("NFKD", clean(value)).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def category_for(description):
    text = normalized(description)
    groups = [
        ("Mangueiras", ("MANGUEIRA", "MANG ", "TUBO PU", "TUBO PNEUM")),
        ("Conexões e adaptadores", ("CONEXAO", "ADAPTADOR", "ADAPT ", "JOELHO", "CURVA ", "TE ", "LUVA ", "BUCHA ", "NIPLE", "ENGATE")),
        ("Ferramentas", ("ALICATE", "CHAVE ", "MARTELO", "SERRA ", "BROCA", "DISCO ", "FURADEIRA", "PARAFUSADEIRA", "ESMERILHADEIRA")),
        ("Fixação", ("PARAFUSO", "PORCA", "ARRUELA", "ABRACADEIRA", "REBITE", "PREGO ")),
        ("Elétrica", ("CABO ", "FIO ", "TOMADA", "INTERRUPTOR", "EXTENSAO", "PLUGUE", "DISJUNTOR")),
        ("Hidráulica", ("REGISTRO", "VALVULA", "SIFAO", "TORNEIRA", "CAIXA D AGUA", "PVC", "SOLDAVEL", "ESGOTO")),
        ("Lubrificação e químicos", ("OLEO", "GRAXA", "LUBRIFIC", "ADESIVO", "COLA ", "SILICONE", "DESENGRIPANTE")),
        ("EPI e segurança", ("LUVA ", "BOTA ", "OCULOS", "CAPACETE", "MASCARA", "PROTETOR")),
    ]
    for name, terms in groups:
        if any(term in text for term in terms):
            return name
    return "Diversos"


existing = json.loads(TARGET.read_text(encoding="utf-8"))
existing_descriptions = {normalized(item["description"]) for item in existing}
known_codes = {clean(item.get("code")) for item in existing if clean(item.get("code"))}

workbook = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook.active
added = []
seen_codes = set(known_codes)
seen_fallback = {(normalized(item["description"]), normalized(item.get("supplier"))) for item in existing}

for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
    code, description, alias, cost, price, supplier, ncm = (list(row) + [None] * 7)[:7]
    code = clean(code)
    description = clean(description)
    supplier = clean(supplier) or "Fornecedor não informado"
    if not description:
        continue
    fallback = (normalized(description), normalized(supplier))
    if (code and code in seen_codes) or normalized(description) in existing_descriptions or (not code and fallback in seen_fallback):
        continue
    if code:
        seen_codes.add(code)
    seen_fallback.add(fallback)
    added.append({
        "id": f"catalogo-{code}" if code else f"catalogo-linha-{row_number}",
        "supplier": supplier,
        "brand": clean(alias) or supplier,
        "code": code,
        "description": description,
        "category": category_for(description),
        "unit": "un",
        "stock": 0,
        "suggested": 0,
        "image": None,
        "sourceRow": row_number,
        "cost": float(cost) if isinstance(cost, (int, float)) else None,
        "price": float(price) if isinstance(price, (int, float)) else None,
        "ncm": clean(ncm),
    })

TARGET.write_text(json.dumps(existing + added, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(json.dumps({"existing": len(existing), "added": len(added), "total": len(existing) + len(added)}, ensure_ascii=False))
